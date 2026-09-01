"""Mechanical Markdown and DOCX handoff for an ordered Evidence Set."""

from __future__ import annotations

import hashlib
import html
import json
import os
import posixpath
import re
import secrets
import shutil
import stat
import threading
import unicodedata
import zipfile
from collections.abc import Callable, Iterator
from contextlib import contextmanager, suppress
from dataclasses import dataclass
from enum import StrEnum
from io import BytesIO
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree

from PIL import Image
from PIL.PngImagePlugin import PngInfo

from csbox.core.safe_paths import (
    atomic_write_text,
    is_reparse_metadata,
    mkdir_exclusive,
    open_regular_binary,
    read_regular_text,
    restore_backup_or_preserve,
    safe_relative_path,
    safe_rename,
)
from csbox.evidence.models import EvidenceItem, EvidenceSet
from csbox.evidence.resolver import LabCaptureResolver
from csbox.lab.models import CaptureRecord
from csbox.lab.renderer import TerminalEvidenceRenderer

_COMPONENT_BUDGET = 240
_MANIFEST_NAME = ".csbox-generated-report.json"
_MAX_MANIFEST_BYTES = 1024 * 1024
_MAX_MARKDOWN_BYTES = 32 * 1024 * 1024
_IMAGE_RELATIONSHIP_SUFFIX = "/image"
_A_NAMESPACE = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"


class ReportExportPhase(StrEnum):
    """The externally visible phases of one report export."""

    GENERATING = "generating"
    PUBLISHING = "publishing"
    COMPLETE = "complete"
    FAILED = "failed"


@dataclass(frozen=True, slots=True)
class ReportExportRequest:
    """Destination and explicit overwrite authorization for a report export."""

    destination: Path
    force: bool = False


@dataclass(frozen=True, slots=True)
class ReportExportResult:
    """Final paths and counts for one successfully published report bundle."""

    destination: Path
    markdown: Path
    docx: Path
    images: tuple[Path, ...]
    evidence_item_count: int
    warnings: tuple[str, ...] = ()


class ReportExportError(RuntimeError):
    """A report cannot be generated or published without an unsafe partial result."""

    def __init__(
        self,
        message: str,
        *,
        phase: ReportExportPhase = ReportExportPhase.FAILED,
        unavailable_titles: tuple[str, ...] = (),
        destination: Path | None = None,
        kind: str = "export",
    ) -> None:
        super().__init__(message)
        self.phase = phase
        self.unavailable_titles = unavailable_titles
        self.destination = destination
        self.kind = kind


PhaseCallback = Callable[[ReportExportPhase], None]
DocxWriter = Callable[[EvidenceSet, tuple[tuple[EvidenceItem, Path], ...], Path], None]


@dataclass(frozen=True, slots=True)
class _ManifestState:
    valid: bool
    entries: dict[PurePosixPath, str]
    existing_targets: frozenset[PurePosixPath] = frozenset()


@dataclass(frozen=True, slots=True)
class _Backup:
    destination: Path
    backup: Path


_FileIdentity = tuple[int, int, int, int, int]


@dataclass(frozen=True, slots=True)
class _CreatedTarget:
    destination: Path
    expected_digest: str
    identity: _FileIdentity | None


_DESTINATION_LOCKS: dict[str, threading.Lock] = {}
_DESTINATION_LOCKS_GUARD = threading.Lock()


@contextmanager
def _destination_lock(destination: Path) -> Iterator[None]:
    """Serialize cooperating publishers for one canonical destination."""

    key = os.path.normcase(os.path.abspath(destination))
    with _DESTINATION_LOCKS_GUARD:
        lock = _DESTINATION_LOCKS.setdefault(key, threading.Lock())
    if not lock.acquire(blocking=False):
        raise ReportExportError(
            "该导出位置正在使用，请稍后重试。",
            destination=destination,
            kind="busy",
        )
    try:
        yield
    finally:
        lock.release()


class EvidenceReportExporter:
    """Export an Evidence Set through canonical Lab Captures only."""

    def __init__(
        self,
        resolver: LabCaptureResolver,
        renderer: TerminalEvidenceRenderer,
        *,
        docx_writer: DocxWriter | None = None,
    ) -> None:
        self.resolver = resolver
        self.renderer = renderer
        self.docx_writer = docx_writer or _write_docx

    def export(
        self,
        evidence_set: EvidenceSet,
        destination: Path | str,
        *,
        force: bool = False,
        phase_callback: PhaseCallback | None = None,
    ) -> ReportExportResult:
        if not isinstance(evidence_set, EvidenceSet):
            raise TypeError("evidence_set must be an EvidenceSet")

        output = Path(destination)
        self._emit(phase_callback, ReportExportPhase.GENERATING)
        staging: Path | None = None
        try:
            resolved = self._resolve_all(evidence_set)
            self._reject_source_destination(output)
            with _destination_lock(output):
                output_existed = _path_exists(output)
                manifest = _preflight_output(
                    output,
                    force=force,
                    expected_names=tuple(
                        _asset_name(index, item.title)
                        for index, item in enumerate(evidence_set.items, start=1)
                    ),
                )
                staging = _new_staging(output)
                assets = staging / "assets"
                mkdir_exclusive(assets)

                rendered: list[tuple[EvidenceItem, Path]] = []
                for index, (item, capture) in enumerate(resolved, start=1):
                    image_path = assets / _asset_name(index, item.title)
                    self.renderer.render(capture.snapshot, image_path)
                    rendered.append((item, image_path))

                markdown_path = staging / "report.md"
                _write_markdown(markdown_path, evidence_set, rendered)
                docx_path = staging / "report.docx"
                self.docx_writer(evidence_set, tuple(rendered), docx_path)
                manifest_path = staging / _MANIFEST_NAME
                _write_manifest(manifest_path, rendered, staging)
                expected_names = tuple(path.name for _item, path in rendered)
                _validate_bundle(staging, expected_names)

                self._emit(phase_callback, ReportExportPhase.PUBLISHING)
                warnings = _publish_bundle(
                    staging,
                    output,
                    output_existed=output_existed,
                    owned_assets=manifest.entries if manifest.valid else {},
                    existing_targets=manifest.existing_targets,
                    expected_names=expected_names,
                )
                warnings = (*warnings, *_cleanup_staging(staging))
                result = ReportExportResult(
                    destination=output,
                    markdown=output / "report.md",
                    docx=output / "report.docx",
                    images=tuple(output / "assets" / name for name in expected_names),
                    evidence_item_count=len(rendered),
                    warnings=warnings,
                )
                self._emit(phase_callback, ReportExportPhase.COMPLETE)
                return result
        except ReportExportError:
            if staging is not None:
                _cleanup_staging(staging)
            self._emit(phase_callback, ReportExportPhase.FAILED)
            raise
        except Exception as exc:
            if staging is not None:
                _cleanup_staging(staging)
            self._emit(phase_callback, ReportExportPhase.FAILED)
            raise ReportExportError(
                _controlled_message(exc),
                destination=output,
                kind=_error_kind(exc),
            ) from exc

    def _resolve_all(
        self,
        evidence_set: EvidenceSet,
    ) -> tuple[tuple[EvidenceItem, CaptureRecord], ...]:
        resolved: list[tuple[EvidenceItem, CaptureRecord]] = []
        unavailable: list[str] = []
        for item in evidence_set.items:
            try:
                source = self.resolver.resolve(item.source)
            except Exception:
                source = None
            if source is None or not source.available or source.capture is None:
                unavailable.append(item.title)
                continue
            resolved.append((item, source.capture))
        if unavailable:
            raise ReportExportError(
                "有证据来源不可用，请恢复来源、移除不可用引用，或返回 Evidence Set。",
                unavailable_titles=tuple(unavailable),
                kind="missing_source",
            )
        return tuple(resolved)

    def _reject_source_destination(self, destination: Path) -> None:
        repository = getattr(self.resolver, "repository", None)
        source_root = getattr(repository, "root", None)
        if source_root is None:
            return
        candidate = Path(os.path.abspath(destination))
        source_root = Path(os.path.abspath(Path(source_root)))
        if candidate == source_root or candidate.is_relative_to(source_root):
            raise ReportExportError(
                "报告不能写入 Lab 源目录，请选择源目录之外的位置。",
                destination=destination,
                kind="destination",
            )

    @staticmethod
    def _emit(callback: PhaseCallback | None, phase: ReportExportPhase) -> None:
        if callback is not None:
            with suppress(Exception):
                callback(phase)


def _controlled_message(error: Exception) -> str:
    if isinstance(error, PermissionError):
        return "报告材料写入失败，请检查目标目录权限后重试。"
    if isinstance(error, FileNotFoundError):
        return "报告材料导出位置不可用，请检查目录后重试。"
    if isinstance(error, (ValueError, NotADirectoryError, IsADirectoryError)):
        return "报告材料导出位置无效，请选择一个可写目录后重试。"
    return "报告材料导出失败，未生成成功结果；请检查目标目录后重试。"


def _error_kind(error: Exception) -> str:
    if isinstance(error, PermissionError):
        return "permission"
    if isinstance(error, (ValueError, NotADirectoryError, IsADirectoryError)):
        return "destination"
    return "generation"


def _resolve_asset_component(name: str) -> PurePosixPath:
    relative = safe_relative_path(name)
    if (
        relative.as_posix() != name
        or len(relative.parts) != 2
        or relative.parts[0] != "assets"
        or not _safe_asset_component(relative.parts[1])
    ):
        raise ValueError("invalid report asset path")
    return relative


def _safe_asset_component(component: str) -> bool:
    if (
        not component
        or component in {".", ".."}
        or component != component.rstrip(" .")
        or any(unicodedata.category(character).startswith("C") for character in component)
        or any(character in '/\\:*?"<>|' for character in component)
        or not component.lower().endswith(".png")
        or len(component.encode("utf-8")) > _COMPONENT_BUDGET
        or _windows_units(component) > _COMPONENT_BUDGET
    ):
        return False
    stem = component.split(".", maxsplit=1)[0].upper()
    return stem not in {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        *(f"COM{index}" for index in range(1, 10)),
        *(f"LPT{index}" for index in range(1, 10)),
    }


def _path_metadata(path: Path) -> os.stat_result | None:
    try:
        return path.lstat()
    except FileNotFoundError:
        return None


def _path_exists(path: Path) -> bool:
    return _path_metadata(path) is not None


def _is_regular(path: Path) -> bool:
    metadata = _path_metadata(path)
    return (
        metadata is not None
        and not is_reparse_metadata(metadata)
        and stat.S_ISREG(metadata.st_mode)
    )


def _is_directory(path: Path) -> bool:
    metadata = _path_metadata(path)
    return (
        metadata is not None
        and not is_reparse_metadata(metadata)
        and stat.S_ISDIR(metadata.st_mode)
    )


def _reject_parent_links(path: Path) -> None:
    absolute = Path(os.path.abspath(path))
    current = Path(absolute.anchor)
    for component in absolute.parts[1:]:
        current /= component
        metadata = _path_metadata(current)
        if metadata is None:
            break
        if is_reparse_metadata(metadata):
            raise ValueError("refusing to write through a symlink parent")
        if not stat.S_ISDIR(metadata.st_mode):
            raise NotADirectoryError(current)


def _preflight_output(
    output: Path,
    *,
    force: bool,
    expected_names: tuple[str, ...],
) -> _ManifestState:
    _reject_parent_links(output.parent)
    metadata = _path_metadata(output)
    if metadata is not None and is_reparse_metadata(metadata):
        raise ReportExportError(
            "导出路径不能是符号链接或重解析点。",
            destination=output,
            kind="destination",
        )
    if metadata is not None and not stat.S_ISDIR(metadata.st_mode):
        raise ReportExportError(
            "导出路径不是目录，请选择一个目录。",
            destination=output,
            kind="destination",
        )
    if metadata is None:
        return _ManifestState(valid=False, entries={})
    if not force:
        raise ReportExportError(
            "导出目录已存在；请确认覆盖后重试。",
            destination=output,
            kind="existing_target",
        )

    assets = output / "assets"
    assets_metadata = _path_metadata(assets)
    if assets_metadata is not None and (
        is_reparse_metadata(assets_metadata) or not stat.S_ISDIR(assets_metadata.st_mode)
    ):
        raise ReportExportError(
            "现有 assets 目录不可安全使用，请选择其他目录。",
            destination=output,
            kind="destination",
        )
    for path in (
        output / "report.md",
        output / "report.docx",
        output / _MANIFEST_NAME,
    ):
        if _path_exists(path) and not _is_regular(path):
            raise ReportExportError(
                "现有报告目标不是普通文件，无法安全覆盖。",
                destination=output,
                kind="destination",
            )

    manifest = _read_manifest(output)
    existing_targets: set[PurePosixPath] = set()
    for path in (
        output / "report.md",
        output / "report.docx",
        output / _MANIFEST_NAME,
    ):
        if _path_exists(path):
            existing_targets.add(PurePosixPath(path.name))
    for name in expected_names:
        path = assets / name
        if not _path_exists(path):
            continue
        if not _is_regular(path):
            raise ReportExportError(
                "现有报告图片不是普通文件，无法安全覆盖。",
                destination=output,
                kind="destination",
            )
        relative = PurePosixPath("assets", name)
        if not manifest.valid or relative not in manifest.entries:
            raise ReportExportError(
                "现有报告图片未被 CSBox 标记为可覆盖，请选择其他目录。",
                destination=output,
                kind="existing_target",
            )
        existing_targets.add(PurePosixPath("assets", name))
    for relative in manifest.entries:
        if _path_exists(output.joinpath(*relative.parts)):
            existing_targets.add(relative)
    return _ManifestState(
        valid=manifest.valid,
        entries=manifest.entries,
        existing_targets=frozenset(existing_targets),
    )


def _new_staging(output: Path) -> Path:
    _reject_parent_links(output.parent)
    for _ in range(16):
        staging = output.parent / f".csbox-report-{secrets.token_hex(16)}.partial"
        try:
            mkdir_exclusive(staging)
        except FileExistsError:
            continue
        return staging
    raise ReportExportError(
        "无法创建安全的报告临时目录，请稍后重试。",
        destination=output,
        kind="destination",
    )


def _asset_name(index: int, title: str) -> str:
    return f"{index:02d}-{_safe_slug(title, index)}.png"


def _safe_slug(title: str, index: int) -> str:
    normalized = unicodedata.normalize("NFKC", title)
    characters: list[str] = []
    for character in normalized:
        if unicodedata.category(character).startswith("C"):
            continue
        if character in '<>:"/\\|?*' or character.isspace():
            characters.append("-")
        else:
            characters.append(character)
    slug = re.sub(r"-+", "-", "".join(characters)).strip(" .-_")
    while ".." in slug:
        slug = slug.replace("..", ".")
    if not slug:
        slug = f"实验记录-{index}"
    if slug.split(".", maxsplit=1)[0].upper() in {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        *(f"COM{number}" for number in range(1, 10)),
        *(f"LPT{number}" for number in range(1, 10)),
    }:
        slug = f"_{slug}"

    prefix = f"{index:02d}-"
    extension = ".png"
    byte_budget = _COMPONENT_BUDGET - len((prefix + extension).encode("utf-8"))
    windows_budget = _COMPONENT_BUDGET - _windows_units(prefix + extension)
    if len(slug.encode("utf-8")) > byte_budget or _windows_units(slug) > windows_budget:
        digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:8]
        suffix = f"-{digest}"
        byte_budget -= len(suffix.encode("utf-8"))
        windows_budget -= _windows_units(suffix)
        shortened: list[str] = []
        for character in slug:
            candidate = "".join(shortened) + character
            if (
                len(candidate.encode("utf-8")) > byte_budget
                or _windows_units(candidate) > windows_budget
            ):
                break
            shortened.append(character)
        slug = "".join(shortened).rstrip(" .-_") + suffix
    return slug


def _windows_units(value: str) -> int:
    return len(value.encode("utf-16-le")) // 2


def _escape_markdown_text(value: str) -> str:
    return re.sub(r"([\\\x60*_{}\[\]()#+|%])", r"\\\1", value)


def _markdown_user_text(value: str) -> str:
    if "\n" not in value and "\r" not in value:
        return _escape_markdown_text(value)
    normalized = value.replace("\r\n", "\n").replace("\r", "\n")
    escaped = _escape_markdown_text(html.escape(normalized, quote=False)).replace("\n", "<br>")
    return "<span>" + escaped + "</span>"


def _markdown_url(value: str) -> str:
    encoded: list[str] = []
    for character in value:
        if ord(character) > 127 or character.isalnum() or character in "-._~/":
            encoded.append(character)
        else:
            encoded.extend(f"%{byte:02X}" for byte in character.encode("utf-8"))
    return "".join(encoded)


def _write_markdown(
    destination: Path,
    evidence_set: EvidenceSet,
    rendered: list[tuple[EvidenceItem, Path]],
) -> None:
    parts = [f"# {_markdown_user_text(evidence_set.title)}\n"]
    for index, (item, image_path) in enumerate(rendered, start=1):
        caption = item.caption if item.caption != "" else item.title
        relative = image_path.relative_to(destination.parent).as_posix()
        parts.append(
            f"\n## {_markdown_user_text(item.title)}\n\n"
            f"![图 {index}]({_markdown_url(relative)})\n\n"
            f"图 {index} {_markdown_user_text(caption)}\n"
        )
        if item.note != "":
            parts.append(f"\n{_markdown_user_text(item.note)}\n")
    atomic_write_text(destination, "".join(parts))


def _write_manifest(
    destination: Path,
    rendered: list[tuple[EvidenceItem, Path]],
    root: Path,
) -> None:
    files = [
        {
            "path": image.relative_to(root).as_posix(),
            "sha256": _file_sha256(image),
        }
        for _item, image in rendered
    ]
    atomic_write_text(
        destination,
        json.dumps({"version": 1, "files": files}, ensure_ascii=False, indent=2) + "\n",
    )


def _read_manifest(root: Path) -> _ManifestState:
    path = root / _MANIFEST_NAME
    if not _path_exists(path) or not _is_regular(path):
        return _ManifestState(valid=False, entries={})
    try:
        raw = json.loads(read_regular_text(path, max_bytes=_MAX_MANIFEST_BYTES))
        if (
            not isinstance(raw, dict)
            or set(raw) != {"version", "files"}
            or type(raw["version"]) is not int
            or raw["version"] != 1
            or not isinstance(raw["files"], list)
        ):
            return _ManifestState(valid=False, entries={})
        entries: dict[PurePosixPath, str] = {}
        for item in raw["files"]:
            if (
                not isinstance(item, dict)
                or set(item) != {"path", "sha256"}
                or not isinstance(item["path"], str)
                or not isinstance(item["sha256"], str)
                or not re.fullmatch(r"[0-9a-f]{64}", item["sha256"])
            ):
                return _ManifestState(valid=False, entries={})
            relative = _resolve_asset_component(item["path"])
            if relative in entries:
                return _ManifestState(valid=False, entries={})
            entries[relative] = item["sha256"]
        for relative, expected in entries.items():
            path = root.joinpath(*relative.parts)
            if not _path_exists(path):
                continue
            if not _is_regular(path) or _file_sha256(path) != expected:
                return _ManifestState(valid=False, entries={})
        return _ManifestState(valid=True, entries=entries)
    except (OSError, UnicodeError, ValueError, json.JSONDecodeError, RecursionError):
        return _ManifestState(valid=False, entries={})


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with open_regular_binary(path) as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _validate_bundle(root: Path, expected_names: tuple[str, ...]) -> None:
    assets = root / "assets"
    if not _is_directory(assets):
        raise ValueError("report assets directory is invalid")
    for name in expected_names:
        path = assets / name
        if not _is_regular(path):
            raise ValueError("report image is invalid")
        _validate_png(path)

    markdown = root / "report.md"
    if not _is_regular(markdown):
        raise ValueError("report Markdown is invalid")
    markdown_text = read_regular_text(markdown, max_bytes=_MAX_MARKDOWN_BYTES)
    if "\x00" in markdown_text:
        raise ValueError("report Markdown contains an invalid character")

    docx = root / "report.docx"
    if not _is_regular(docx):
        raise ValueError("report DOCX is invalid")
    _validate_docx(docx, len(expected_names))

    manifest = _read_manifest(root)
    expected_entries = {
        PurePosixPath("assets", name): _file_sha256(root / "assets" / name)
        for name in expected_names
    }
    if not manifest.valid or manifest.entries != expected_entries:
        raise ValueError("report manifest is invalid")


def _validate_png(path: Path) -> None:
    try:
        with Image.open(path) as image:
            if image.format != "PNG":
                raise ValueError("report image is not a PNG")
            image.verify()
    except Exception as exc:
        raise ValueError("report image is corrupt") from exc


def _validate_docx(path: Path, expected_image_count: int) -> None:
    try:
        with zipfile.ZipFile(path) as package:
            if package.testzip() is not None:
                raise ValueError("DOCX contains a corrupt member")
            names = set(package.namelist())
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ValueError("DOCX package is incomplete")
            media = sorted(
                name for name in names if name.startswith("word/media/") and not name.endswith("/")
            )
            if len(media) != expected_image_count:
                raise ValueError("DOCX image count is incorrect")
            if any(name.startswith("word/fonts/") for name in names):
                raise ValueError("DOCX must not embed fonts")
            document = ElementTree.fromstring(package.read("word/document.xml"))
            image_blips = document.findall(f".//{{{_A_NAMESPACE}}}blip")
            if len(image_blips) != expected_image_count:
                raise ValueError("DOCX image references are incorrect")
            document_relationships: dict[str, ElementTree.Element] = {}
            for name in names:
                if not name.endswith(".rels"):
                    continue
                relationships = ElementTree.fromstring(package.read(name))
                for relationship in relationships:
                    if relationship.tag != f"{{{_RELATIONSHIP_NAMESPACE}}}Relationship":
                        continue
                    if (
                        relationship.attrib.get("Type", "").endswith(_IMAGE_RELATIONSHIP_SUFFIX)
                        and relationship.attrib.get("TargetMode", "").lower() == "external"
                    ):
                        raise ValueError("DOCX contains an external image relationship")
                    if name == "word/_rels/document.xml.rels":
                        relationship_id = relationship.attrib.get("Id")
                        if relationship_id is None or relationship_id in document_relationships:
                            raise ValueError("DOCX document relationships are invalid")
                        document_relationships[relationship_id] = relationship

            media_set = set(media)
            referenced_media: set[str] = set()
            embed_attribute = f"{{{_R_NAMESPACE}}}embed"
            link_attribute = f"{{{_R_NAMESPACE}}}link"
            for image_blip in image_blips:
                if link_attribute in image_blip.attrib:
                    raise ValueError("DOCX contains a linked image")
                relationship_id = image_blip.attrib.get(embed_attribute)
                relationship = document_relationships.get(relationship_id or "")
                if relationship is None or not relationship.attrib.get("Type", "").endswith(
                    _IMAGE_RELATIONSHIP_SUFFIX
                ):
                    raise ValueError("DOCX image relationship is missing")
                if relationship.attrib.get("TargetMode", "").lower() == "external":
                    raise ValueError("DOCX contains an external image relationship")
                target = relationship.attrib.get("Target", "")
                part_name = _docx_relationship_part(target)
                if part_name not in media_set:
                    raise ValueError("DOCX image relationship target is invalid")
                referenced_media.add(part_name)
            if referenced_media != media_set:
                raise ValueError("DOCX media is not fully embedded and referenced")
            for name in media:
                with Image.open(BytesIO(package.read(name))) as image:
                    if image.format != "PNG":
                        raise ValueError("DOCX embedded image is not a PNG")
                    image.verify()
    except (OSError, ValueError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise ValueError("DOCX package is invalid") from exc


def _docx_relationship_part(target: str) -> str:
    if not target or target.startswith("/"):
        raise ValueError("DOCX relationship target is not a package-relative path")
    normalized = PurePosixPath(posixpath.normpath((PurePosixPath("word") / target).as_posix()))
    if not normalized.parts or normalized.parts[0] != "word":
        raise ValueError("DOCX relationship target escapes the word package")
    return normalized.as_posix()


def _write_docx(
    evidence_set: EvidenceSet,
    rendered: tuple[tuple[EvidenceItem, Path], ...],
    destination: Path,
) -> None:
    """Create the fixed linear handoff document with embedded rendered images."""

    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.core_properties.title = evidence_set.title
    document.add_heading(evidence_set.title, level=0)
    section = document.sections[0]
    available_width = section.page_width - section.left_margin - section.right_margin
    if available_width <= 0:
        available_width = Inches(6)

    temporary_images: list[Path] = []
    try:
        for index, (item, image_path) in enumerate(rendered, start=1):
            document.add_heading(item.title, level=1)
            temporary_image = destination.parent / (
                f".csbox-docx-image-{index}-{secrets.token_hex(8)}.png"
            )
            _copy_png_with_unique_bytes(image_path, temporary_image, index)
            temporary_images.append(temporary_image)
            document.add_picture(str(temporary_image), width=available_width)

            caption = item.caption if item.caption != "" else item.title
            caption_paragraph = document.add_paragraph()
            _add_preserved_text(caption_paragraph, f"图 {index} {caption}")
            if item.note != "":
                note_paragraph = document.add_paragraph()
                _add_preserved_text(note_paragraph, item.note)
        document.save(destination)
    finally:
        for temporary_image in temporary_images:
            with suppress(OSError):
                temporary_image.unlink()


def _copy_png_with_unique_bytes(source: Path, destination: Path, index: int) -> None:
    """Give each DOCX picture distinct PNG bytes so every item embeds one image."""

    with Image.open(source) as image:
        metadata = PngInfo()
        metadata.add_text("csbox-report-item", str(index))
        image.save(destination, format="PNG", pnginfo=metadata)


def _add_preserved_text(paragraph: object, text: str) -> None:
    """Write user text while representing its existing line boundaries as breaks."""

    parts = re.split(r"(\r\n|\r|\n)", text)
    for part in parts:
        if part in {"\r\n", "\r", "\n"}:
            paragraph.add_run().add_break()  # type: ignore[attr-defined]
        else:
            paragraph.add_run(part)  # type: ignore[attr-defined]


def _publish_bundle(
    staging: Path,
    output: Path,
    *,
    output_existed: bool,
    owned_assets: dict[PurePosixPath, str],
    existing_targets: frozenset[PurePosixPath],
    expected_names: tuple[str, ...],
) -> tuple[str, ...]:
    """Publish a validated bundle through a rollback-capable per-file transaction."""

    created_directories: list[Path] = []
    created_targets: list[_CreatedTarget] = []
    backups: list[_Backup] = []
    warnings: list[str] = []

    try:
        if not output_existed:
            if _path_exists(output):
                raise ValueError("report destination changed before publication")
            mkdir_exclusive(output)
            created_directories.append(output)
        elif not _is_directory(output):
            raise ValueError("report destination changed before publication")
        assets = output / "assets"
        if not _is_directory(assets):
            if _path_exists(assets):
                raise ValueError("report assets destination changed before publication")
            mkdir_exclusive(assets)
            created_directories.append(assets)

        staged_images = tuple(staging / "assets" / name for name in expected_names)
        generated = tuple((source, output / "assets" / source.name) for source in staged_images) + (
            (staging / "report.md", output / "report.md"),
            (staging / "report.docx", output / "report.docx"),
            (staging / _MANIFEST_NAME, output / _MANIFEST_NAME),
        )
        generated_digests = {destination: _file_sha256(source) for source, destination in generated}

        for source, destination in generated:
            relative = _relative_path(destination, output)
            expected_owner = owned_assets.get(relative)
            if relative not in existing_targets:
                unowned_asset = relative.parts and relative.parts[0] == "assets"
                if _path_exists(destination):
                    raise ReportExportError(
                        (
                            "发布前检测到未授权的报告图片冲突，请选择其他目录。"
                            if unowned_asset
                            else "发布前检测到目标文件冲突，请选择其他目录。"
                        ),
                        destination=output,
                        kind="existing_target",
                    )
                try:
                    _publish_created_target(
                        source,
                        destination,
                        generated_digests[destination],
                        created_targets,
                    )
                except FileExistsError as exc:
                    raise ReportExportError(
                        "发布时检测到目标文件冲突，请选择其他目录。",
                        destination=output,
                        kind="existing_target",
                    ) from exc
                except ValueError as exc:
                    raise ReportExportError(
                        "报告目标不安全，未覆盖用户内容；请选择其他目录。",
                        destination=output,
                        kind="existing_target",
                    ) from exc
                continue
            backup = _move_existing_to_backup(
                destination,
                staging,
                len(backups) + 1,
                expected_owner=expected_owner
                if relative.parts and relative.parts[0] == "assets"
                else None,
                report_destination=output,
            )
            if backup is not None:
                backups.append(backup)
            _publish_created_target(
                source,
                destination,
                generated_digests[destination],
                created_targets,
            )

        for relative, expected_owner in owned_assets.items():
            if relative in {PurePosixPath("assets", name) for name in expected_names}:
                continue
            if relative not in existing_targets:
                continue
            destination = output.joinpath(*relative.parts)
            if not _path_exists(destination):
                continue
            backup = _move_existing_to_backup(
                destination,
                staging,
                len(backups) + 1,
                expected_owner=expected_owner,
                report_destination=output,
            )
            if backup is not None:
                backups.append(backup)

        _validate_bundle(output, expected_names)
    except Exception as exc:
        rollback_errors = _rollback_publication(
            created_targets,
            backups,
            created_directories,
        )
        if rollback_errors:
            raise ReportExportError(
                "报告发布失败，部分并发冲突已保留为恢复文件；请检查后重试。",
                destination=output,
                kind="recovery",
            ) from exc
        if isinstance(exc, ReportExportError):
            raise
        raise ReportExportError(
            "报告发布失败，旧的有效报告已保留；请重试。",
            destination=output,
            kind="publication",
        ) from exc

    for backup in backups:
        try:
            backup.backup.unlink()
        except OSError:
            warnings.append("旧报告备份未能清理，已保留恢复文件；当前报告仍可用。")
    return tuple(warnings)


def _relative_path(path: Path, root: Path) -> PurePosixPath:
    return PurePosixPath(path.relative_to(root).as_posix())


def _file_identity(metadata: os.stat_result | None) -> _FileIdentity | None:
    if metadata is None:
        return None
    inode = getattr(metadata, "st_ino", 0)
    if not inode:
        return None
    return (
        int(getattr(metadata, "st_dev", 0)),
        int(inode),
        int(getattr(metadata, "st_size", -1)),
        int(getattr(metadata, "st_mtime_ns", 0)),
        int(getattr(metadata, "st_ctime_ns", 0)),
    )


def _same_file_location(left: _FileIdentity, right: _FileIdentity) -> bool:
    return left[:2] == right[:2]


def _publish_created_target(
    source: Path,
    destination: Path,
    expected_digest: str,
    created_targets: list[_CreatedTarget],
) -> None:
    source_identity = _file_identity(_path_metadata(source))
    safe_rename(source, destination, replace_existing=False)
    destination_identity = _file_identity(_path_metadata(destination))
    if destination_identity is None:
        created_targets.append(_CreatedTarget(destination, expected_digest, None))
        return
    if source_identity is not None and not _same_file_location(
        source_identity,
        destination_identity,
    ):
        created_targets.append(_CreatedTarget(destination, expected_digest, None))
        raise ValueError("published report target changed during publication")
    created_targets.append(_CreatedTarget(destination, expected_digest, destination_identity))


def _move_existing_to_backup(
    destination: Path,
    staging: Path,
    number: int,
    *,
    expected_owner: str | None,
    report_destination: Path,
) -> _Backup | None:
    metadata = _path_metadata(destination)
    if metadata is None:
        return None
    if is_reparse_metadata(metadata) or not stat.S_ISREG(metadata.st_mode):
        raise ValueError("report target is not a regular file")
    if expected_owner is not None and _file_sha256(destination) != expected_owner:
        raise ReportExportError(
            "现有报告图片已被修改，未覆盖用户内容；请确认后重试。",
            destination=report_destination,
            kind="existing_target",
        )

    backup = staging / f".backup-{number}"
    safe_rename(destination, backup, replace_existing=False)
    if expected_owner is not None:
        try:
            verified = _file_sha256(backup)
        except (OSError, ValueError) as exc:
            raise ReportExportError(
                "现有报告图片无法安全验证，未覆盖用户内容。",
                destination=report_destination,
                kind="existing_target",
            ) from exc
        if verified != expected_owner:
            _restore_single_backup(backup, destination)
            raise ReportExportError(
                "现有报告图片在备份时发生变化，未覆盖用户内容。",
                destination=report_destination,
                kind="existing_target",
            )
    return _Backup(destination=destination, backup=backup)


def _restore_single_backup(backup: Path, destination: Path) -> None:
    restored = restore_backup_or_preserve(backup, destination)
    if restored != destination:
        raise ReportExportError(
            "旧报告文件无法安全恢复；恢复文件已保留，请检查后重试。",
            destination=destination,
            kind="recovery",
        )


def _rollback_publication(
    created_targets: list[_CreatedTarget],
    backups: list[_Backup],
    created_directories: list[Path],
) -> tuple[str, ...]:
    errors: list[str] = []
    for created_target in reversed(created_targets):
        try:
            _remove_created_target(created_target)
        except Exception:
            errors.append(f"无法处理已创建目标：{created_target.destination.name}")
    for backup in reversed(backups):
        if not _path_exists(backup.backup):
            continue
        try:
            _restore_single_backup(backup.backup, backup.destination)
        except Exception:
            errors.append(f"无法恢复旧目标：{backup.destination.name}")
    for directory in reversed(created_directories):
        try:
            metadata = _path_metadata(directory)
            if (
                metadata is not None
                and not is_reparse_metadata(metadata)
                and stat.S_ISDIR(metadata.st_mode)
                and not any(directory.iterdir())
            ):
                directory.rmdir()
        except Exception:
            errors.append(f"无法清理临时目录：{directory.name}")
    return tuple(errors)


def _remove_created_target(created_target: _CreatedTarget) -> None:
    destination = created_target.destination
    metadata = _path_metadata(destination)
    if metadata is None:
        return
    if is_reparse_metadata(metadata) or not stat.S_ISREG(metadata.st_mode):
        raise ValueError("created report target is no longer a regular file")
    current_identity = _file_identity(_path_metadata(destination))
    if (
        created_target.identity is None
        or current_identity is None
        or current_identity != created_target.identity
    ):
        raise ValueError("created report target was replaced")
    if _file_sha256(destination) != created_target.expected_digest:
        raise ValueError("created report target content changed")
    recovery = _recovery_path(destination)
    safe_rename(destination, recovery, replace_existing=False)
    recovery_metadata = _path_metadata(recovery)
    recovery_identity = _file_identity(recovery_metadata)
    recovery_digest: str | None
    try:
        recovery_digest = _file_sha256(recovery)
    except (OSError, ValueError):
        recovery_digest = None
    if (
        recovery_identity is None
        or not _same_file_location(recovery_identity, created_target.identity)
        or recovery_digest != created_target.expected_digest
    ):
        with suppress(FileExistsError, OSError, ValueError):
            safe_rename(recovery, destination, replace_existing=False)
        raise ValueError("created report target changed during rollback")
    recovery.unlink()


def _recovery_path(destination: Path) -> Path:
    return destination.parent / f".csbox-recovery-{secrets.token_hex(16)}.bak"


def _cleanup_staging(staging: Path) -> tuple[str, ...]:
    if not _path_exists(staging):
        return ()
    if not _is_directory(staging):
        return ("报告临时路径未能清理，请检查目标目录。",)
    try:
        if any(path.name.startswith(".backup-") for path in staging.iterdir()):
            return ("旧报告备份未能清理，已保留恢复文件；请检查目标目录。",)
        shutil.rmtree(staging)
    except OSError:
        return ("报告临时目录未能清理，请检查目标目录。",)
    return () if not _path_exists(staging) else ("报告临时目录未能清理，请检查目标目录。",)


__all__ = [
    "DocxWriter",
    "EvidenceReportExporter",
    "PhaseCallback",
    "ReportExportError",
    "ReportExportPhase",
    "ReportExportRequest",
    "ReportExportResult",
]
